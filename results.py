import pandas as pd
import re
from pathlib import Path
from openpyxl import load_workbook, Workbook

def write_results(results_by_job, workbook_path):

    workbook_path = Path(workbook_path)

    # Get or create workbook
    if workbook_path.exists():
        wb = load_workbook(workbook_path)

    else:
        wb = Workbook()

        # Openpyxl adds blank sheet by default
        if "Sheet" in wb.sheetnames:
            wb.remove(wb["Sheet"])


    def append_files_to_sheet(paths, workbook, sheet_name, job_name):

        for path in paths:
            # reads space-seperated text into pandas df
            df = pd.read_csv(path, sep=r"\s+")

            # Adding a job label for readability in Excel file
            match = re.search(r"_model_(\d+)", path.name)
            model_index = match.group(1) if match else "?"
            job_label = f"{job_name}_{model_index}"

            # Get or create worksheet
            if sheet_name in workbook.sheetnames:
                worksheet = workbook[sheet_name]

            else:
                worksheet = workbook.create_sheet(sheet_name)
                header = ["job_name"] + df.columns.tolist()
                worksheet.append(header)

            # Writing each row of the file and adding the job_name to differentiate row data
            for data in df.values.tolist():
                row = [job_label] + data
                worksheet.append(row)

    # Loop file to their Excel sheet
    for job_name, (chain_score_paths, residue_score_paths) in results_by_job.items():
        append_files_to_sheet(chain_score_paths, wb, "Output chain-chain score", job_name)
        append_files_to_sheet(residue_score_paths, wb, "Output by-residue score", job_name)

    wb.save(workbook_path)